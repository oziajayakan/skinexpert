"""
Script deterministik untuk mengekstrak data produk dari dataset/DataProduk.xlsx
dan menyimpannya dalam format JSON untuk digunakan oleh sistem pakar web.
"""

import json
import os
import openpyxl

def extract_products():
    excel_path = os.path.join(os.path.dirname(__file__), "..", "dataset", "DataProduk.xlsx")
    output_path = os.path.join(os.path.dirname(__file__), "..", "dataset", "produk_rekomendasi.json")

    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active

    headers = [cell for cell in next(sheet.iter_rows(values_only=True))]
    products = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        prod = {
            "no": row[0],
            "nama": row[1],
            "merek": row[2],
            "kategori": row[3],
            "bpom": row[4],
            "spf_pa": row[5],
            "tekstur": row[6],
            "harga": row[7],
            "tipe_kulit": [s.strip() for s in str(row[8]).split(",") if s.strip()],
            "target_masalah": [s.strip() for s in str(row[9]).split(",") if s.strip()],
            "kandungan": [line.strip().lstrip("-•* ") for line in str(row[10]).split("\n") if line.strip()],
            "fungsi": [line.strip().lstrip("-•* ") for line in str(row[11]).split("\n") if line.strip()]
        }
        products.append(prod)

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(products, f, indent=2, ensure_ascii=False)

    print(f"Berhasil mengekstrak {len(products)} produk ke {output_path}")

if __name__ == "__main__":
    extract_products()
